#!/usr/bin/env python3
"""Extract workplace/life cases from 12 docx files and build 3 trilingual HTML pages."""
import docx
import os
import re
from html import escape

SRC_DIR = "/Users/tai7hermes/YiwordSite_MVP/docs/source-docs/AI-易經-A11-64掛與384爻-案例-3-202609/"
OUT_DIR = "/Users/tai7hermes/YiwordSite_MVP/"

FILES = {
    'A31': 'A31-周易六十四卦_職場生活決策_事件故事與卦義對照_繁體中文_網站App_RWD工程版.docx',
    'A32': 'A32-周易六十四卦_歐美職場生活決策_事件故事與卦義對照_繁體中文_網站App_RWD工程版.docx',
    'B31': 'B31-易傳小象384爻_中華職場生活決策_事件故事與歷史對照_繁體中文_網站App_RWD工程版.docx',
    'B32': 'B32-易傳小象384爻_歐美職場生活決策_事件故事與歷史對照_繁體中文_網站App_RWD工程版.docx',
    'C31': 'C31-周易六十四卦_中华职场与生活决策史案例_简体中文版_网站App_RWD.docx',
    'C32': 'C32-周易六十四卦_欧美职场与生活决策史案例_简体中文版_网站App_RWD.docx',
    'D31': 'D31-易传小象384爻_职场与生活决策_简体中文版.docx',
    'D32': 'D32-易传小象384爻_欧美职场与生活决策_简体中文版.docx',
    'E31': 'E31_Zhouyi_64_Workplace_Life_Historical_Stories_English_Web_App_RWD_Expanded_Draft.docx',
    'E32': 'E32-Zhouyi_64_Workplace_Life_Decisions_English_REISSUED.docx',
    'F31': 'F31-Yizhuan_XiaoXiang_H01-H64_English_Workplace_Life_Decisions_001-384_Complete_Web_App_RWD.docx',
    'F32': 'F32-Yizhuan_WL_001_384_Workplace_Life_Decisions_COMPLETE.docx',
}


def load_doc(label):
    return docx.Document(os.path.join(SRC_DIR, FILES[label]))


def get_case_indices(doc, heading_style_name, pattern):
    """Get paragraph indices where the heading matches the pattern."""
    result = []
    for i, p in enumerate(doc.paragraphs):
        if p.style and heading_style_name in p.style.name:
            t = p.text.strip()
            if pattern(t):
                result.append(i)
    return result


def extract_normal_paras(doc, start, end):
    """Extract non-empty Normal paragraphs between start+1 and end."""
    paras = []
    for j in range(start + 1, end):
        p = doc.paragraphs[j]
        if p.style and 'Normal' in p.style.name:
            t = p.text.strip()
            if t:
                paras.append(t)
    return paras


def find_next_heading(doc, start, heading_styles):
    """Find the index of the next heading after start that matches any of heading_styles."""
    for j in range(start + 1, len(doc.paragraphs)):
        p = doc.paragraphs[j]
        if p.style:
            for hs in heading_styles:
                if hs in p.style.name:
                    return j
    return len(doc.paragraphs)


def extract_64_cases(doc, heading_style):
    """Extract 64卦 cases. heading_style = 'Heading 1' or 'Heading 2'.
    Returns list of (heading, paras[]) where paras includes all Normal text."""
    pattern = lambda t: re.match(r'^\d{2}\s', t) is not None
    indices = get_case_indices(doc, heading_style, pattern)
    cases = []
    for k, idx in enumerate(indices):
        next_idx = indices[k + 1] if k + 1 < len(indices) else len(doc.paragraphs)
        heading = doc.paragraphs[idx].text.strip()
        paras = extract_normal_paras(doc, idx, next_idx)
        cases.append((heading, paras))
    return cases


def extract_384_cases_h2_numbered(doc):
    """For B31, D31: H2 headings 'NN 卦 爻'."""
    pattern = lambda t: re.match(r'^\d{2}\s', t) is not None
    indices = get_case_indices(doc, 'Heading 2', pattern)
    cases = []
    for k, idx in enumerate(indices):
        next_idx = indices[k + 1] if k + 1 < len(indices) else len(doc.paragraphs)
        heading = doc.paragraphs[idx].text.strip()
        paras = extract_normal_paras(doc, idx, next_idx)
        cases.append((heading, paras))
    return cases


def extract_384_cases_h2_yao(doc):
    """For B32, D32: H2 headings '第N爻｜...'."""
    pattern = lambda t: re.match(r'^第\d+爻', t) is not None
    indices = get_case_indices(doc, 'Heading 2', pattern)
    cases = []
    for k, idx in enumerate(indices):
        next_idx = indices[k + 1] if k + 1 < len(indices) else len(doc.paragraphs)
        heading = doc.paragraphs[idx].text.strip()
        paras = extract_normal_paras(doc, idx, next_idx)
        cases.append((heading, paras))
    return cases


def extract_384_cases_h3_numbered(doc):
    """For F31: H3 headings 'NNN 初九 | ...'."""
    pattern = lambda t: re.match(r'^\d{3}\s', t) is not None
    indices = get_case_indices(doc, 'Heading 3', pattern)
    cases = []
    for k, idx in enumerate(indices):
        next_idx = indices[k + 1] if k + 1 < len(indices) else len(doc.paragraphs)
        heading = doc.paragraphs[idx].text.strip()
        paras = extract_normal_paras(doc, idx, next_idx)
        cases.append((heading, paras))
    return cases


def extract_384_cases_f32(doc):
    """For F32: mix of H3 'Line N' (sections 1-3) and H2 'Line N' (section 4).
    Also H2 hexagram headings and H3 line headings within sections."""
    # Get all paragraphs that are headings starting with 'Line '
    line_indices = []
    for i, p in enumerate(doc.paragraphs):
        if p.style:
            t = p.text.strip()
            if t.startswith('Line ') and ('Heading 2' in p.style.name or 'Heading 3' in p.style.name):
                line_indices.append(i)

    cases = []
    for k, idx in enumerate(line_indices):
        next_idx = line_indices[k + 1] if k + 1 < len(line_indices) else len(doc.paragraphs)
        heading = doc.paragraphs[idx].text.strip()
        paras = extract_normal_paras(doc, idx, next_idx)
        cases.append((heading, paras))
    return cases


# ---- Parse case content into structured format ----

def parse_64_zh(heading, paras):
    """Parse 繁中/簡中 64卦 case (A31/A32/C31/C32).
    paras[0] = meta line (ID ｜ date ｜ era)
    paras[1] = tag line (B｜... or A｜...)
    paras[2:] = body paragraphs
    """
    meta = paras[0] if paras else ''
    tag = ''
    body = []
    if len(paras) > 1:
        m = re.match(r'^([AB])', paras[1])
        if m:
            tag = m.group(1)
            body = paras[2:]
        else:
            body = paras[1:]
    else:
        body = paras[1:]
    return heading, meta, tag, body


def parse_64_en_e31(heading, paras):
    """Parse E31 English 64卦 case.
    paras[0] = 'LIFE-H01-001  |  H01-J  |  Type B — ...'
    """
    meta = paras[0] if paras else ''
    tag = ''
    body = paras[1:] if len(paras) > 1 else []
    m = re.search(r'Type\s+([AB])', meta)
    if m:
        tag = m.group(1)
    return heading, meta, tag, body


def parse_64_en_e32(heading, paras):
    """Parse E32 English 歐美 64卦 case.
    paras[0] = 'EUUS-WL-H01  ·  Rome, Europe · 458 BCE  ·  Evidence: RETROSPECTIVE ANALOGY'
    """
    meta = paras[0] if paras else ''
    tag = 'B'  # All E32 cases are RETROSPECTIVE ANALOGY
    body = paras[1:] if len(paras) > 1 else []
    return heading, meta, tag, body


def parse_384_zh(heading, paras, line_num):
    """Parse 繁中/簡中 384爻 case (B31/D31).
    paras[0] = 'XLIFE-H01-L1　｜　XIANG-U-01-02　｜　B 事後對照...'
    """
    meta = paras[0] if paras else ''
    tag = f'{line_num:03d}爻'
    body = paras[1:] if len(paras) > 1 else []
    return heading, meta, tag, body


def parse_384_b32(heading, paras, line_num):
    """Parse B32 繁中 歐美 384爻 case.
    paras[0] = 'source_unit_id：XIANG-U-01-02｜line_no：1'
    """
    meta = paras[0] if paras else ''
    tag = f'{line_num:03d}爻'
    body = paras[1:] if len(paras) > 1 else []
    return heading, meta, tag, body


def parse_384_d32(heading, paras, line_num):
    """Parse D32 簡中 歐美 384爻 case.
    paras[0] = 'source_unit_id: XIANG-U-01-02   |   event_id: ...'
    """
    meta = paras[0] if paras else ''
    tag = f'{line_num:03d}爻'
    body = paras[1:] if len(paras) > 1 else []
    return heading, meta, tag, body


def parse_384_en_f31(heading, paras, line_num):
    """Parse F31 English 384爻 case.
    heading = '001  初九  |  Build your skills...'
    paras[0] = 'line_id: XLIFE-H01-L1    xiang_id: ...'
    """
    meta = paras[0] if paras else ''
    tag = f'{line_num:03d}爻'
    body = paras[1:] if len(paras) > 1 else []
    return heading, meta, tag, body


def parse_384_en_f32(heading, paras, line_num):
    """Parse F32 English 歐美 384爻 case.
    heading = 'Line 1  |  TW-WL-H01-L1  |  XIANG-U-01-02'
    """
    meta = paras[0] if paras else ''
    tag = f'{line_num:03d}爻'
    body = paras[1:] if len(paras) > 1 else []
    return heading, meta, tag, body


# ---- HTML generation ----

def make_case_entry(heading, meta, tag, body):
    """Generate a <div class="case-entry"> HTML block."""
    parts = []
    parts.append('<div class="case-entry">')
    parts.append(f'  <h4>{escape(heading)}</h4>')
    if meta or tag:
        meta_escaped = escape(meta) if meta else ''
        tag_escaped = escape(tag) if tag else ''
        parts.append(f'  <div class="case-meta"><span class="case-tag">{tag_escaped}</span> {meta_escaped}</div>')
    for p in body:
        parts.append(f'  <p>{escape(p)}</p>')
    parts.append('</div>')
    return '\n'.join(parts)


def build_html(lang, label_text, section1_title, section2_title,
               sub1_title, sub2_title,
               zh_hua_64, zh_hua_384, ou_hua_64, ou_hua_384):
    """Build a complete HTML page.
    lang = 'zh-Hant', 'zh-Hans', or 'en'
    """
    if lang == 'zh-Hant':
        # Root directory
        css_href = 'style.css'
        lang_switch = ('<div class="lang-switch">'
                       '<a href="cn/case-workplace.html">简中</a> '
                       '<a href="case-workplace.html" class="active">繁中</a> '
                       '<a href="en/case-workplace.html">EN</a>'
                       '</div>')
        nav_links = [
            ('index.html', '首頁'),
            ('introduction.html', '易經介紹'),
            ('hexagrams.html', '六十四卦'),
            ('wings.html', '易傳十翼'),
            ('cases.html', '案例'),
            ('divination.html', '占卜'),
            ('terms.html', '名詞索引'),
            ('search.html', '搜尋'),
            ('about.html', '關於'),
        ]
        back_home_href = 'index.html'
        back_home_text = '⌂ 首頁'
        back_text = '← 回上頁'
        footer_priv = 'privacy.html'
        footer_priv_text = '隱私權政策'
        footer_terms = 'terms.html'
        footer_terms_text = '使用條款'
        brand_text = '易理 = 處世之道'
    elif lang == 'zh-Hans':
        css_href = '../style.css'
        lang_switch = ('<div class="lang-switch">'
                       '<a href="case-workplace.html" class="active">简中</a> '
                       '<a href="../case-workplace.html">繁中</a> '
                       '<a href="../en/case-workplace.html">EN</a>'
                       '</div>')
        nav_links = [
            ('index.html', '首页'),
            ('introduction.html', '易经介绍'),
            ('hexagrams.html', '六十四卦'),
            ('wings.html', '易传十翼'),
            ('cases.html', '案例'),
            ('divination.html', '占卜'),
            ('terms.html', '名词索引'),
            ('search.html', '搜索'),
            ('about.html', '关于'),
        ]
        back_home_href = 'index.html'
        back_home_text = '⌂ 首页'
        back_text = '← 回上页'
        footer_priv = 'privacy.html'
        footer_priv_text = '隐私政策'
        footer_terms = 'terms.html'
        footer_terms_text = '使用条款'
        brand_text = '易理 = 处世之道'
    else:  # en
        css_href = '../style.css'
        lang_switch = ('<div class="lang-switch">'
                       '<a href="../cn/case-workplace.html">简中</a> '
                       '<a href="../case-workplace.html">繁中</a> '
                       '<a href="case-workplace.html" class="active">EN</a>'
                       '</div>')
        nav_links = [
            ('index.html', 'Home'),
            ('introduction.html', 'Introduction'),
            ('hexagrams.html', '64 Hexagrams'),
            ('wings.html', 'Ten Wings'),
            ('cases.html', 'Cases'),
            ('divination.html', 'Divination'),
            ('terms.html', 'Glossary'),
            ('search.html', 'Search'),
            ('about.html', 'About'),
        ]
        back_home_href = 'index.html'
        back_home_text = '⌂ Home'
        back_text = '← Back'
        footer_priv = 'privacy.html'
        footer_priv_text = 'Privacy Policy'
        footer_terms = 'terms.html'
        footer_terms_text = 'Terms of Use'
        brand_text = 'Yi = Word of Modus Operandi'

    # Build nav
    nav_items = []
    for href, text in nav_links:
        active = ' class="active"' if text in ('案例', 'Cases') else ''
        nav_items.append(f'<li><a href="{href}"{active}>{text}</a></li>')
    nav_html = '<nav><ul>' + ''.join(nav_items) + '</ul></nav>'

    # Build case sections
    def build_cases_html(cases):
        return '\n'.join([make_case_entry(h, m, t, b) for h, m, t, b in cases])

    s1_64 = build_cases_html(zh_hua_64)
    s1_384 = build_cases_html(zh_hua_384)
    s2_64 = build_cases_html(ou_hua_64)
    s2_384 = build_cases_html(ou_hua_384)

    html = f'''<!DOCTYPE html>
<html lang="{lang}">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>{escape(label_text)} — YiWord 易理</title>
<link rel="stylesheet" href="{css_href}">
</head>
<body>
<header><div class="brand">YiWord <span>{escape(brand_text)}</span></div>{lang_switch}</header>
{nav_html}
<main>
<div class="back-nav"><a href="{back_home_href}">{back_home_text}</a><a href="javascript:history.back()">{back_text}</a></div>
<div class="hex-num-label">{escape(label_text)}</div>

<h2>{escape(section1_title)}</h2>
<h3>{escape(sub1_title)}</h3>
{s1_64}
<h3>{escape(sub2_title)}</h3>
{s1_384}

<h2>{escape(section2_title)}</h2>
<h3>{escape(sub1_title)}</h3>
{s2_64}
<h3>{escape(sub2_title)}</h3>
{s2_384}

</main>
<footer><p>© 2026 YiWord.com · <a href="{footer_priv}">{footer_priv_text}</a> · <a href="{footer_terms}">{footer_terms_text}</a></p></footer>
</body>
</html>'''
    return html


def main():
    # ---- 繁中 (A31, A32, B31, B32) ----
    # A31: 中華 64卦
    doc_a31 = load_doc('A31')
    a31_raw = extract_64_cases(doc_a31, 'Heading 2')
    a31_cases = [parse_64_zh(h, p) for h, p in a31_raw]

    # A32: 歐美 64卦
    doc_a32 = load_doc('A32')
    a32_raw = extract_64_cases(doc_a32, 'Heading 1')
    a32_cases = [parse_64_zh(h, p) for h, p in a32_raw]

    # B31: 中華 384爻
    doc_b31 = load_doc('B31')
    b31_raw = extract_384_cases_h2_numbered(doc_b31)
    b31_cases = [parse_384_zh(h, p, i + 1) for i, (h, p) in enumerate(b31_raw)]

    # B32: 歐美 384爻
    doc_b32 = load_doc('B32')
    b32_raw = extract_384_cases_h2_yao(doc_b32)
    b32_cases = [parse_384_b32(h, p, i + 1) for i, (h, p) in enumerate(b32_raw)]

    # ---- 簡中 (C31, C32, D31, D32) ----
    # C31: 中華 64卦
    doc_c31 = load_doc('C31')
    c31_raw = extract_64_cases(doc_c31, 'Heading 2')
    c31_cases = [parse_64_zh(h, p) for h, p in c31_raw]

    # C32: 歐美 64卦
    doc_c32 = load_doc('C32')
    c32_raw = extract_64_cases(doc_c32, 'Heading 1')
    c32_cases = [parse_64_zh(h, p) for h, p in c32_raw]

    # D31: 中華 384爻
    doc_d31 = load_doc('D31')
    d31_raw = extract_384_cases_h2_numbered(doc_d31)
    d31_cases = [parse_384_zh(h, p, i + 1) for i, (h, p) in enumerate(d31_raw)]

    # D32: 歐美 384爻
    doc_d32 = load_doc('D32')
    d32_raw = extract_384_cases_h2_yao(doc_d32)
    d32_cases = [parse_384_d32(h, p, i + 1) for i, (h, p) in enumerate(d32_raw)]

    # ---- English (E31, E32, F31, F32) ----
    # E31: 中華 64卦
    doc_e31 = load_doc('E31')
    e31_raw = extract_64_cases(doc_e31, 'Heading 1')
    e31_cases = [parse_64_en_e31(h, p) for h, p in e31_raw]

    # E32: 歐美 64卦
    doc_e32 = load_doc('E32')
    e32_raw = extract_64_cases(doc_e32, 'Heading 2')
    e32_cases = [parse_64_en_e32(h, p) for h, p in e32_raw]

    # F31: 中華 384爻
    doc_f31 = load_doc('F31')
    f31_raw = extract_384_cases_h3_numbered(doc_f31)
    f31_cases = [parse_384_en_f31(h, p, i + 1) for i, (h, p) in enumerate(f31_raw)]

    # F32: 歐美 384爻
    doc_f32 = load_doc('F32')
    f32_raw = extract_384_cases_f32(doc_f32)
    f32_cases = [parse_384_en_f32(h, p, i + 1) for i, (h, p) in enumerate(f32_raw)]

    # Print counts
    counts = {
        'A31': len(a31_cases), 'A32': len(a32_cases),
        'B31': len(b31_cases), 'B32': len(b32_cases),
        'C31': len(c31_cases), 'C32': len(c32_cases),
        'D31': len(d31_cases), 'D32': len(d32_cases),
        'E31': len(e31_cases), 'E32': len(e32_cases),
        'F31': len(f31_cases), 'F32': len(f32_cases),
    }
    for k, v in counts.items():
        print(f"  {k}: {v} cases")

    total = sum(counts.values())
    print(f"Total cases: {total}")

    # Build 繁中 HTML
    tw_html = build_html(
        'zh-Hant', '職場與生活案例',
        '中華案例', '歐美案例',
        '六十四卦案例', '三百八十四爻案例',
        a31_cases, b31_cases, a32_cases, b32_cases,
    )
    tw_path = os.path.join(OUT_DIR, 'case-workplace.html')
    with open(tw_path, 'w', encoding='utf-8') as f:
        f.write(tw_html)
    print(f"\nWrote: {tw_path} ({len(tw_html)} bytes)")

    # Build 簡中 HTML
    cn_html = build_html(
        'zh-Hans', '职场与生活案例',
        '中华案例', '欧美案例',
        '六十四卦案例', '三百八十四爻案例',
        c31_cases, d31_cases, c32_cases, d32_cases,
    )
    cn_path = os.path.join(OUT_DIR, 'cn', 'case-workplace.html')
    with open(cn_path, 'w', encoding='utf-8') as f:
        f.write(cn_html)
    print(f"Wrote: {cn_path} ({len(cn_html)} bytes)")

    # Build English HTML
    en_html = build_html(
        'en', 'Workplace and Life Cases',
        'Chinese Cases', 'Western Cases',
        '64 Hexagram Cases', '384 Line Cases',
        e31_cases, f31_cases, e32_cases, f32_cases,
    )
    en_path = os.path.join(OUT_DIR, 'en', 'case-workplace.html')
    with open(en_path, 'w', encoding='utf-8') as f:
        f.write(en_html)
    print(f"Wrote: {en_path} ({len(en_html)} bytes)")

    return total


if __name__ == '__main__':
    total = main()
    print(f"\n=== DONE: {total} total cases extracted ===")
